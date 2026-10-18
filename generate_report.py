import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from sqlalchemy import create_engine, text

# 1. KẾT NỐI DATABASE SQLITE
engine = create_engine('sqlite:///tdic_data.db')

# 2. THỰC HIỆN CÁC TRUY VẤN SQL PHÂN TÍCH TÀI CHÍNH
query_summary = """
SELECT 
    COUNT(DISTINCT Invoice_No) AS Total_Invoices,
    SUM(Pre_Tax_Total) AS Total_Pre_Tax,
    SUM(Tax_Amount) AS Total_VAT,
    SUM(Grand_Total) AS Total_Grand
FROM invoices;
"""

query_vendor = """
SELECT 
    Vendor_Name,
    Vendor_MST,
    COUNT(Invoice_No) AS Invoice_Count,
    SUM(Pre_Tax_Total) AS Total_Pre_Tax,
    SUM(Tax_Amount) AS Total_VAT,
    SUM(Grand_Total) AS Total_Grand,
    ROUND(SUM(Grand_Total) * 100.0 / (SELECT SUM(Grand_Total) FROM invoices), 2) AS Cost_Share_Pct
FROM invoices
GROUP BY Vendor_Name, Vendor_MST
ORDER BY Total_Grand DESC;
"""

query_top_items = """
SELECT 
    Item_Name,
    Unit,
    SUM(Quantity) AS Total_Quantity,
    AVG(Unit_Price) AS Avg_Unit_Price,
    SUM(Line_Total) AS Total_Spent
FROM invoice_details
GROUP BY Item_Name, Unit
ORDER BY Total_Spent DESC
LIMIT 15;
"""

query_tax_declaration = """
SELECT 
    Invoice_No,
    Vendor_Name,
    Vendor_MST,
    Pre_Tax_Total,
    Tax_Rate,
    Tax_Amount,
    Grand_Total
FROM invoices
ORDER BY Invoice_No ASC;
"""

# Đọc dữ liệu vào Pandas DataFrame
df_summary = pd.read_sql(query_summary, con=engine)
df_vendor = pd.read_sql(query_vendor, con=engine)
df_top_items = pd.read_sql(query_top_items, con=engine)
df_tax = pd.read_sql(query_tax_declaration, con=engine)

# 3. XUẤT VÀ ĐỊNH DẠNG FILE EXCEL CHUYÊN NGHIỆP VỚI OPENPYXL
output_file = 'TDIC_Financial_Report.xlsx'
wb = Workbook()
# Xóa sheet mặc định
wb.remove(wb.active)

# Định dạng Style
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_header = PatternFill(
    start_color='1F4E78', end_color='1F4E78', fill_type='solid'
)
font_title = Font(name='Calibri', size=14, bold=True, color='1F4E78')
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)


def apply_custom_formatting(
    ws, df, title, number_cols=[], currency_cols=[], pct_cols=[]
):
  # Thêm tiêu đề Sheet
  ws.append([title])
  ws.cell(row=1, column=1).font = font_title
  ws.append([])  # Dòng trống

  # Thêm dữ liệu DataFrame
  for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

  # Formatting cho Header Row (Dòng 3)
  for col in range(1, len(df.columns) + 1):
    cell = ws.cell(row=3, column=col)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(
        horizontal='center', vertical='center', wrap_text=True
    )

  # Formatting cho Data Rows
  for row in range(4, 4 + len(df)):
    for col in range(1, len(df.columns) + 1):
      cell = ws.cell(row=row, column=col)
      cell.border = border_thin

      col_name = df.columns[col - 1]
      if col_name in currency_cols:
        cell.number_format = '#,##0" VNĐ"'
        cell.alignment = Alignment(horizontal='right')
      elif col_name in pct_cols:
        cell.number_format = '0.00"%"'
        cell.alignment = Alignment(horizontal='right')
      elif col_name in number_cols:
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal='right')
      else:
        cell.alignment = Alignment(horizontal='left')

  # Auto-fit độ rộng cột
  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


# --- SHEET 1: EXECUTIVE SUMMARY ---
ws1 = wb.create_sheet(title='Tổng Quan Chi Phí')
apply_custom_formatting(
    ws1,
    df_summary,
    'BÁO CÁO TỔNG QUAN CHI PHÍ HOÁ ĐƠN',
    number_cols=['Total_Invoices'],
    currency_cols=['Total_Pre_Tax', 'Total_VAT', 'Total_Grand'],
)

# --- SHEET 2: VENDOR ANALYSIS ---
ws2 = wb.create_sheet(title='Phân Tích Nhà Cung Cấp')
apply_custom_formatting(
    ws2,
    df_vendor,
    'PHÂN TÍCH CHI PHÍ THEO NHÀ CUNG CẤP',
    number_cols=['Invoice_Count'],
    currency_cols=['Total_Pre_Tax', 'Total_VAT', 'Total_Grand'],
    pct_cols=['Cost_Share_Pct'],
)

# --- SHEET 3: TOP COST ITEMS ---
ws3 = wb.create_sheet(title='Top Vật Tư Chi Phí Cao')
apply_custom_formatting(
    ws3,
    df_top_items,
    'TOP 15 VẬT TƯ / MẶT HÀNG CHI PHÍ CAO NHẤT',
    number_cols=['Total_Quantity'],
    currency_cols=['Avg_Unit_Price', 'Total_Spent'],
)

# --- SHEET 4: TAX DECLARATION ---
ws4 = wb.create_sheet(title='Bảng Kê Thuế VAT Input')
apply_custom_formatting(
    ws4,
    df_tax,
    'BẢNG KÊ THUẾ GTGT ĐẦU VÀO KÊ KHAI',
    currency_cols=['Pre_Tax_Total', 'Tax_Amount', 'Grand_Total'],
    pct_cols=['Tax_Rate'],
)

# Lưu Workbook
wb.save(output_file)
print(f' SUCCESS: Đã tạo báo cáo tài chính tự động tại `{output_file}`!')